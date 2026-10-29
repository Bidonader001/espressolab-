"""Incomplete days and days off must be impossible to miss in every report."""

from __future__ import annotations

import csv
import io

import pytest
from openpyxl import load_workbook

from test_app import _device
from test_market_ready import _load

PARAMS = {"report": "monthly", "start_date": "2026-10-01", "end_date": "2026-10-04", "source": "device"}


@pytest.fixture()
def person(client, auth):
    device = _device(client, auth, name="HQ")
    logs = [
        {"user_id": "20", "punched_at": "2026-10-01T09:00:00", "punch": 0},  # Thursday: check-in, no check-out
        {"user_id": "20", "punched_at": "2026-10-03T09:00:00", "punch": 0},  # Saturday: worked on a day off
        {"user_id": "20", "punched_at": "2026-10-03T17:00:00", "punch": 1},
    ]
    _load(client, auth, device["id"], [{"user_id": "20", "name": "Iman Incomplete"}], logs)


def _csv(client, auth) -> dict[str, list[str]]:
    text = client.get("/api/v1/exports", headers=auth, params={**PARAMS, "format": "csv"}).content.decode("utf-8-sig")
    return {row[2][:10]: row for row in csv.reader(io.StringIO(text)) if row[0] == "20"}


def test_incomplete_day_spells_out_the_missing_check_out(client, auth, person):
    thursday = _csv(client, auth)["01/10/2026"]
    assert thursday[3] == "09:00" and thursday[4] == "No check-out" and thursday[5] == "" and thursday[6] == "Incomplete"


def test_friday_and_saturday_are_listed_and_never_counted_absent(client, auth, person):
    rows = _csv(client, auth)
    assert rows["02/10/2026"][6] == "Weekend" and rows["02/10/2026"][2].endswith("Fri")
    assert rows["02/10/2026"][3:6] == ["", "", ""]
    # Saturday has punches, so it is a normal worked row rather than a Weekend placeholder.
    assert rows["03/10/2026"][6] in {"Present", "Late"} and rows["03/10/2026"][5] == "8:00"
    summary = next(
        row for row in client.get("/api/v1/attendance", headers=auth, params={**PARAMS, "page_size": 200}).json()["summary"]
        if row["employee_id"] == "20"
    )
    # Thursday is incomplete, Sunday 04/10 is a scheduled day with no punches, Friday and Saturday are days off.
    assert summary["incomplete_days"] == 1 and summary["absences"] == 1


def test_api_marks_days_off_so_screens_can_highlight_them(client, auth, person):
    items = client.get("/api/v1/attendance", headers=auth, params={**PARAMS, "page_size": 200}).json()["items"]
    mine = {row["date_iso"]: row for row in items if row["employee_id"] == "20"}
    assert mine["2026-10-02"]["status"] == "Weekend" and mine["2026-10-02"]["weekend"] is True
    assert mine["2026-10-03"]["weekend"] is True and mine["2026-10-03"]["status"] != "Weekend"
    assert mine["2026-10-01"]["weekend"] is False and mine["2026-10-01"]["status"] == "Incomplete"
    only = client.get("/api/v1/attendance", headers=auth, params={**PARAMS, "status": "Weekend", "page_size": 200}).json()
    assert only["items"] and {row["status"] for row in only["items"]} == {"Weekend"}


def test_workbook_highlights_incomplete_rows_and_days_off(client, auth, person):
    content = client.get("/api/v1/exports", headers=auth, params={**PARAMS, "format": "xlsx"}).content
    sheet = load_workbook(io.BytesIO(content))["Attendance Details"]
    rows = {str(sheet.cell(r, 3).value)[:10]: r for r in range(9, sheet.max_row + 1) if sheet.cell(r, 1).value == "20"}

    def fill(day, col):
        return sheet.cell(rows[day], col).fill.fgColor.rgb

    assert [fill("01/10/2026", col) for col in range(1, 8)] == ["00FFE08A"] * 7  # Incomplete: whole row
    assert sheet.cell(rows["01/10/2026"], 5).value == "No check-out"
    assert [fill("02/10/2026", col) for col in range(1, 8)] == ["00D6E4F5"] * 7  # Friday: whole row
    assert fill("03/10/2026", 3) == "00D6E4F5"  # worked Saturday: the date cell


def test_pdf_still_builds_with_highlighted_rows(client, auth, person):
    pdf = client.get("/api/v1/exports", headers=auth, params={**PARAMS, "format": "pdf"})
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")


def test_working_day_without_any_punch_is_listed_as_absent(client, auth, person):
    """No check-in and no check-out on a working day: the day is a visible Absent row, never skipped."""
    sunday = _csv(client, auth)["04/10/2026"]
    assert sunday[2].endswith("Sun") and sunday[3:6] == ["", "", ""] and sunday[6] == "Absent"
    items = client.get("/api/v1/attendance", headers=auth, params={**PARAMS, "page_size": 200}).json()
    mine = {row["date_iso"]: row for row in items["items"] if row["employee_id"] == "20"}
    assert mine["2026-10-04"]["status"] == "Absent" and mine["2026-10-04"]["weekend"] is False
    # The Absent rows in the list are exactly the absences the summary counts, so the two can never disagree.
    listed = sum(1 for row in items["items"] if row["status"] == "Absent" and row["employee_id"] == "20")
    summary = next(row for row in items["summary"] if row["employee_id"] == "20")
    assert listed == summary["absences"] == 1
    only = client.get("/api/v1/attendance", headers=auth, params={**PARAMS, "status": "Absent", "page_size": 200}).json()
    assert only["items"] and {row["status"] for row in only["items"]} == {"Absent"}
    # Days off and days worked (even incomplete) are never reported as absent.
    assert mine["2026-10-02"]["status"] == "Weekend" and mine["2026-10-01"]["status"] == "Incomplete"
