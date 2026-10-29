"""Attendance rules, exports, device sync, and permissions."""

from __future__ import annotations

import io
from datetime import datetime, timedelta
from types import SimpleNamespace

from openpyxl import load_workbook

from app.time_utils import hours_between


def _hours(start: str, end: str) -> float:
    return round(hours_between(datetime.fromisoformat(start), datetime.fromisoformat(end)), 4)


def _login(client, username, password):
    response = client.post("/api/v1/auth/login", json={"username": username, "password": password})
    assert response.status_code == 200, response.text
    return {"Authorization": f"Bearer {response.json()['token']}"}


def _ids(client, auth):
    branches = client.get("/api/v1/branches", headers=auth).json()["items"]
    shifts = client.get("/api/v1/shifts", headers=auth).json()["items"]
    hq = next(row["id"] for row in branches if row["name"] == "HQ")
    day = next(row["id"] for row in shifts if row["name"] == "Day")
    night = next(row["id"] for row in shifts if row["name"] == "Night")
    return hq, day, night


def _employee(client, auth, employee_id, name, branch_id, shift_id, department="Ops", status="active"):
    response = client.post(
        "/api/v1/employees",
        headers=auth,
        json={
            "employee_id": employee_id,
            "full_name": name,
            "department": department,
            "branch_id": branch_id,
            "shift_id": shift_id,
            "employment_status": status,
        },
    )
    assert response.status_code == 200, response.text
    return response.json()


def _device(client, auth, name="Front door", host="lab-sim.local", protocol="simulator"):
    response = client.post(
        "/api/v1/devices",
        headers=auth,
        json={"name": name, "host": host, "port": 4370, "protocol": protocol, "branch_id": None, "sync_enabled": False},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert "comm_key" not in body
    assert body["has_comm_key"] is False
    return body


def test_health_and_login(client):
    assert client.get("/api/v1/health").json()["ai"] is False
    denied = client.get("/api/v1/employees")
    assert denied.status_code == 401
    bad = client.post("/api/v1/auth/login", json={"username": "admin", "password": "nope"})
    assert bad.status_code == 401


def test_viewer_cannot_write(client, auth):
    created = client.post(
        "/api/v1/users",
        headers=auth,
        json={"username": "viewer1", "password": "viewer-pass-1", "role": "viewer"},
    )
    assert created.status_code == 200, created.text
    viewer = _login(client, "viewer1", "viewer-pass-1")
    blocked = client.post(
        "/api/v1/employees",
        headers=viewer,
        json={"employee_id": "X", "full_name": "Nope"},
    )
    assert blocked.status_code == 403
    assert client.get("/api/v1/attendance", headers=viewer).status_code == 200
    assert client.get("/api/v1/users", headers=viewer).status_code == 403


def test_dashboard_is_empty_until_data_exists(client, auth):
    body = client.get("/api/v1/dashboard", headers=auth).json()
    assert body["present"] == 0
    assert body["absent"] == 0
    assert body["active_employees"] == 0
    assert body["last_sync_at"] is None
    assert body["devices"]["total"] == 0


def test_punch_rules_grace_pairs_overnight_and_duplicates(client, auth):
    hq, day, night = _ids(client, auth)
    _employee(client, auth, "E1", "Ada", hq, day)
    _employee(client, auth, "E2", "Bea", hq, night)
    device = _device(client, auth)
    logs = [
        {"user_id": "10", "punched_at": "2026-10-01T08:10:00", "punch": 0},
        {"user_id": "10", "punched_at": "2026-10-01T16:00:00", "punch": 1},
        {"user_id": "11", "punched_at": "2026-10-01T08:11:00", "punch": 0},
        {"user_id": "11", "punched_at": "2026-10-01T16:00:00", "punch": 1},
        {"user_id": "10", "punched_at": "2026-10-06T08:00:00", "punch": 0},
        {"user_id": "10", "punched_at": "2026-10-06T12:00:00", "punch": 1},
        {"user_id": "10", "punched_at": "2026-10-06T13:00:00", "punch": 0},
        {"user_id": "10", "punched_at": "2026-10-06T17:00:00", "punch": 1},
        {"user_id": "10", "punched_at": "2026-10-05T08:00:00", "punch": 0},
        {"user_id": "12", "punched_at": "2026-10-04T20:00:00", "punch": 0},
        {"user_id": "12", "punched_at": "2026-10-05T04:00:00", "punch": 1},
    ]
    stored = client.post(f"/api/v1/devices/{device['id']}/simulator-logs", headers=auth, json={"logs": logs, "users": []})
    assert stored.json()["simulated"] is True
    assert "ZKTeco" in stored.json()["warning"]
    client.post(f"/api/v1/devices/{device['id']}/map", headers=auth, json={"device_user_id": "10", "employee_id": "E1"})
    client.post(f"/api/v1/devices/{device['id']}/map", headers=auth, json={"device_user_id": "11", "employee_id": "E1"})
    client.post(f"/api/v1/devices/{device['id']}/map", headers=auth, json={"device_user_id": "12", "employee_id": "E2"})
    # User 11 was mapped onto E1 as well, which would merge punches. Keep 11 on a second employee.
    _employee(client, auth, "E3", "Cara", hq, day)
    client.post(f"/api/v1/devices/{device['id']}/map", headers=auth, json={"device_user_id": "11", "employee_id": "E3"})
    first = client.post(f"/api/v1/devices/{device['id']}/sync?wait=1", headers=auth)
    assert first.status_code == 200, first.text
    run = first.json()["run"]
    assert run["status"] == "success"
    assert run["records_inserted"] == len(logs)
    assert first.json()["device"]["connection_status"] == "simulated"
    assert first.json()["device"]["connection_status"] != "online"

    second = client.post(f"/api/v1/devices/{device['id']}/sync?wait=1", headers=auth)
    assert second.json()["run"]["records_inserted"] == 0
    assert second.json()["run"]["records_skipped"] == len(logs)

    rows = client.get(
        "/api/v1/attendance",
        headers=auth,
        params={"start_date": "2026-10-01", "end_date": "2026-10-06", "page_size": 100},
    ).json()["items"]
    by_key = {(row["employee_id"], row["date_iso"]): row for row in rows}
    on_time = by_key[("E1", "2026-10-01")]
    assert on_time["status"] == "Present"
    assert on_time["late_minutes"] == 10
    assert on_time["check_in"] == "08:10"
    late = by_key[("E3", "2026-10-01")]
    assert late["status"] == "Late"
    assert late["late_minutes"] == 11
    lunch = by_key[("E1", "2026-10-06")]
    assert lunch["check_in"] == "08:00"
    assert lunch["check_out"] == "17:00"
    assert lunch["hours"] == 8.0
    missing = by_key[("E1", "2026-10-05")]
    assert missing["status"] == "Incomplete"
    assert missing["check_out"] == ""
    assert missing["hours"] is None
    night = by_key[("E2", "2026-10-04")]
    assert night["check_in"] == "20:00"
    assert night["check_out"] == "04:00"
    assert night["hours"] == 8.0
    # The night session belongs to 10-04. 10-05 has no punch of its own, so it is listed as Absent, not worked.
    assert by_key[("E2", "2026-10-05")]["status"] == "Absent" and by_key[("E2", "2026-10-05")]["check_in"] == ""


def test_offline_sync_keeps_previous_punches(client, auth):
    device = _device(client, auth, host="lab-sim.local")
    client.post(
        f"/api/v1/devices/{device['id']}/simulator-logs",
        headers=auth,
        json={"logs": [{"user_id": "1", "punched_at": "2026-10-01T08:00:00", "punch": 0}], "users": []},
    )
    ok = client.post(f"/api/v1/devices/{device['id']}/sync?wait=1", headers=auth)
    assert ok.json()["run"]["records_inserted"] == 1
    client.patch(f"/api/v1/devices/{device['id']}", headers=auth, json={"host": "offline"})
    failed = client.post(f"/api/v1/devices/{device['id']}/sync?wait=1", headers=auth)
    assert failed.json()["run"]["status"] == "error"
    assert "timed out" in failed.json()["run"]["error_message"].lower()
    punches = client.get(f"/api/v1/devices/{device['id']}/punches", headers=auth).json()["items"]
    assert len(punches) == 1
    assert failed.json()["device"]["connection_status"] == "offline"


def test_push_is_deduped_and_getrequest_is_empty(client, auth):
    created = client.post(
        "/api/v1/devices",
        headers=auth,
        json={
            "name": "Gate",
            "host": "10.0.0.8",
            "protocol": "push",
            "serial_number": "SN100",
            "port": 4370,
        },
    )
    assert created.status_code == 200, created.text
    body = "7\t2026-10-03 09:00:00\t0\t1\t0\t0\n7\t2026-10-03 17:00:00\t1\t1\t0\t0\n"
    first = client.post("/iclock/cdata?SN=SN100&table=ATTLOG", content=body)
    assert first.status_code == 200
    assert first.text == "OK"
    second = client.post("/iclock/cdata?SN=SN100&table=ATTLOG", content=body)
    assert second.text == "OK"
    runs = client.get("/api/v1/sync-runs", headers=auth).json()["items"]
    push_runs = [row for row in runs if row["trigger"] == "push"]
    assert push_runs[0]["records_skipped"] == 2
    assert push_runs[1]["records_inserted"] == 2
    command = client.get("/iclock/getrequest?SN=SN100")
    assert command.text == "OK"
    assert "CLEAR" not in command.text
    unknown = client.get("/iclock/getrequest?SN=MISSING")
    assert unknown.status_code == 404


def test_pyzk_adapter_never_clears(monkeypatch):
    calls = []

    class Conn:
        def get_attendance(self):
            calls.append("attendance")
            return [SimpleNamespace(user_id="4", timestamp=datetime(2026, 10, 1, 8, 0), status=0, punch=0, uid=9)]

        def get_users(self):
            calls.append("users")
            return [SimpleNamespace(user_id="4", name="Ada", uid=9)]

        def get_firmware_version(self):
            return "6.60"

        def get_serialnumber(self):
            return "ABC"

        def get_device_name(self):
            return "K40"

        def disconnect(self):
            calls.append("disconnect")

        def clear_attendance(self):
            calls.append("CLEAR")

        def disable_device(self):
            calls.append("DISABLE")

    class ZK:
        def __init__(self, *args, **kwargs):
            calls.append(kwargs.get("force_udp"))

        def connect(self):
            calls.append("connect")
            return Conn()

    import sys

    monkeypatch.setitem(sys.modules, "zk", SimpleNamespace(ZK=ZK))
    from app.zk.base import ConnParams
    from app.zk import pyzk_adapter

    params = ConnParams(1, "Door", "10.0.0.4", 4370, 5, "tcp", 0, "")
    info = pyzk_adapter.test_connection(params)
    logs, users = pyzk_adapter.fetch(params)
    assert info.ok is True
    assert info.simulated is False
    assert logs[0].user_id == "4"
    assert users[0].name == "Ada"
    assert "CLEAR" not in calls
    assert "DISABLE" not in calls


def test_exports_match_calculation_and_workbook(client, auth):
    hq, day, _night = _ids(client, auth)
    _employee(client, auth, "EMP-001", "Ahmed Mohamed", hq, day)
    _employee(client, auth, "EMP-003", "=2+2", hq, day)
    device = _device(client, auth, name="HQ reader")
    logs = [
        {"user_id": "1", "punched_at": "2026-10-01T08:55:00", "punch": 0},
        {"user_id": "1", "punched_at": "2026-10-01T17:02:00", "punch": 1},
        {"user_id": "3", "punched_at": "2026-10-01T08:00:00", "punch": 0},
        {"user_id": "3", "punched_at": "2026-10-01T16:00:00", "punch": 1},
    ]
    client.post(f"/api/v1/devices/{device['id']}/simulator-logs", headers=auth, json={"logs": logs})
    client.post(f"/api/v1/devices/{device['id']}/map", headers=auth, json={"device_user_id": "1", "employee_id": "EMP-001"})
    client.post(f"/api/v1/devices/{device['id']}/map", headers=auth, json={"device_user_id": "3", "employee_id": "EMP-003"})
    client.post(f"/api/v1/devices/{device['id']}/sync?wait=1", headers=auth)
    params = {"report": "monthly", "format": "xlsx", "start_date": "2026-10-01", "end_date": "2026-10-06", "branch": "HQ"}
    denied = client.get("/api/v1/exports", params=params)
    assert denied.status_code == 401
    bad = client.get("/api/v1/exports", headers=auth, params={**params, "start_date": "2026-10-31", "end_date": "2026-10-01"})
    assert bad.status_code == 400
    workbook = client.get("/api/v1/exports", headers=auth, params=params)
    assert workbook.status_code == 200
    assert "EspressoLabAttendance_Attendance_2026-10.xlsx" in workbook.headers["content-disposition"]
    book = load_workbook(io.BytesIO(workbook.content))
    assert book.sheetnames[0] == "Attendance Details"
    assert "Employee Summary" in book.sheetnames
    sheet = book["Attendance Details"]
    headers = [sheet.cell(8, col).value for col in range(1, 8)]
    assert headers == ["Employee ID", "Name", "Date", "Check-in", "Check-out", "Hours", "Status"]
    assert sheet.freeze_panes == "A9"
    summary = book["Employee Summary"]
    summary_headers = [summary.cell(8, col).value for col in range(1, 7)]
    assert summary_headers == ["Employee", "Days present", "Late arrivals", "Absences", "Incomplete days", "Total hours"]
    assert summary.cell(8, 7).value is None
    # 08:55 to 17:02 is 8 hours 7 minutes. Excel gets a real duration shown as 8:07, never a decimal.
    ahmed_row = next(row for row in range(9, sheet.max_row + 1) if sheet.cell(row, 1).value == "EMP-001")
    assert sheet.cell(ahmed_row, 6).number_format == "[h]:mm"
    assert sheet.cell(ahmed_row, 6).value == timedelta(hours=8, minutes=7)
    csv_file = client.get("/api/v1/exports", headers=auth, params={**params, "format": "csv"})
    text = csv_file.content.decode("utf-8-sig")
    assert text.startswith("Employee ID,")
    assert "'=2+2" in text
    ahmed = next(line for line in text.splitlines() if line.startswith("EMP-001,"))
    assert "08:55" in ahmed
    assert "17:02" in ahmed
    assert ",8:07," in ahmed and "8.12" not in ahmed and "8.1" not in ahmed
    # 08:55 against the 08:00 Day shift is 55 minutes, which is late. The hours still match BunkerAtt.
    assert ahmed.endswith("Late") or ",Late" in ahmed
    pdf = client.get("/api/v1/exports", headers=auth, params={**params, "format": "pdf"})
    assert pdf.content.startswith(b"%PDF")
    assert "Attendance_Summary_2026-10.pdf" in pdf.headers["content-disposition"]
    legacy = client.get("/api/v1/reports/excel", headers=auth, params={"work_date": "2026-10-01", "period": "daily"})
    legacy_book = load_workbook(io.BytesIO(legacy.content))
    assert legacy_book.sheetnames[0] == "Attendance"
    assert legacy_book["Attendance"]["A6"].value == "Name"


def test_leave_and_holiday_are_not_absences(client, auth):
    hq, day, _night = _ids(client, auth)
    _employee(client, auth, "EMP-002", "Mariam Ali", hq, day)
    client.post("/api/v1/holidays", headers=auth, json={"holiday_date": "2026-10-04", "name": "Store holiday", "branch_id": hq})
    client.post("/api/v1/leave", headers=auth, json={"employee_id": "EMP-002", "work_date": "2026-10-05", "note": "Approved"})
    data = client.get(
        "/api/v1/exports/count",
        headers=auth,
        params={"report": "absences", "start_date": "2026-10-01", "end_date": "2026-10-06", "branch": "HQ"},
    ).json()
    # Working days in range: 1 Oct, 5 Oct, 6 Oct. 4 Oct is a holiday, 5 Oct is leave, 2-3 Oct are Friday and Saturday.
    exported = client.get(
        "/api/v1/exports",
        headers=auth,
        params={"report": "absences", "format": "csv", "start_date": "2026-10-01", "end_date": "2026-10-06", "employee_id": "EMP-002"},
    )
    text = exported.content.decode("utf-8-sig")
    assert "04/10/2026" not in text
    assert "05/10/2026" not in text
    assert "02/10/2026" not in text
    assert "01/10/2026" in text
    assert "06/10/2026" in text
    assert data["count"] >= 2


def test_sync_enrolls_device_users_and_export_skips_manual_people(client, auth):
    hq, day, _night = _ids(client, auth)
    _employee(client, auth, "MANUAL-1", "Typed In", hq, day)
    device = _device(client, auth, name="Gate reader")
    stored = client.post(
        f"/api/v1/devices/{device['id']}/simulator-logs",
        headers=auth,
        json={
            "users": [{"user_id": "77", "name": "Device Ada"}],
            "logs": [
                {"user_id": "77", "punched_at": "2026-10-01T08:00:00", "punch": 0},
                {"user_id": "77", "punched_at": "2026-10-01T16:00:00", "punch": 1},
            ],
        },
    )
    assert stored.status_code == 200, stored.text
    synced = client.post(f"/api/v1/devices/{device['id']}/sync?wait=1", headers=auth)
    assert synced.status_code == 200, synced.text
    assert synced.json()["run"]["records_inserted"] == 2
    people = client.get("/api/v1/employees", headers=auth).json()["items"]
    enrolled = next(row for row in people if row["employee_id"] == "77")
    assert enrolled["full_name"] == "Device Ada"
    assert enrolled["origin"] == "device"
    exported = client.get(
        "/api/v1/exports",
        headers=auth,
        params={"report": "daily", "format": "csv", "start_date": "2026-10-01", "end_date": "2026-10-01", "source": "device"},
    )
    text = exported.content.decode("utf-8-sig")
    assert "Device Ada" in text
    assert "08:00" in text
    assert "MANUAL-1" not in text
    assert "Typed In" not in text


def test_deleted_device_attendance_leaves_the_export(client, auth):
    device = _device(client, auth, name="Gate reader")
    client.post(
        f"/api/v1/devices/{device['id']}/simulator-logs",
        headers=auth,
        json={
            "users": [{"user_id": "77", "name": "Device Ada"}],
            "logs": [
                {"user_id": "77", "punched_at": "2026-10-01T08:00:00", "punch": 0},
                {"user_id": "77", "punched_at": "2026-10-01T16:00:00", "punch": 1},
            ],
        },
    )
    synced = client.post(f"/api/v1/devices/{device['id']}/sync?wait=1", headers=auth)
    assert synced.status_code == 200, synced.text
    removed = client.delete(f"/api/v1/devices/{device['id']}", headers=auth)
    assert removed.status_code == 200, removed.text
    assert removed.json()["deleted"] is True
    assert client.get(f"/api/v1/devices/{device['id']}", headers=auth).status_code == 404
    exported = client.get(
        "/api/v1/exports",
        headers=auth,
        params={"report": "daily", "format": "csv", "start_date": "2026-10-01", "end_date": "2026-10-06", "source": "device"},
    )
    text = exported.content.decode("utf-8-sig")
    assert "Device Ada" not in text
    assert "77" not in text
    people = client.get("/api/v1/employees", headers=auth).json()["items"]
    assert all(row["employee_id"] != "77" for row in people)


def test_backup_omits_secrets(client, auth, tmp_path):
    device = _device(client, auth)
    client.patch(f"/api/v1/devices/{device['id']}", headers=auth, json={"comm_key": "12345"})
    saved = client.get(f"/api/v1/devices/{device['id']}", headers=auth).json()
    assert saved["has_comm_key"] is True
    assert "12345" not in str(saved)
    backup = client.post("/api/v1/settings/backup", headers=auth)
    assert backup.status_code == 200, backup.text
    payload = backup.json()["json"]
    text = open(payload, encoding="utf-8").read()
    assert "12345" not in text
    assert "password_hash" not in text


def test_branch_can_be_edited_renamed_and_deleted(client, auth):
    first = client.post("/api/v1/branches", headers=auth, json={"name": "North", "code": "N1"}).json()
    second = client.post("/api/v1/branches", headers=auth, json={"name": "South"}).json()
    edited = client.patch(
        f"/api/v1/branches/{first['id']}",
        headers=auth,
        json={"name": "North Mall", "code": "NM", "timezone": "Europe/Istanbul", "address": "1 Main St", "working_days": "Sun-Thu", "grace_minutes": 5},
    )
    assert edited.status_code == 200, edited.text
    assert edited.json()["name"] == "North Mall" and edited.json()["grace_minutes"] == 5 and edited.json()["working_days"] == "Sun-Thu"
    # Renaming onto another branch's name is refused with a clear message, not a database error.
    clash = client.patch(f"/api/v1/branches/{second['id']}", headers=auth, json={"name": "North Mall"})
    assert clash.status_code == 400 and "already exists" in clash.json()["detail"]
    # Saving a branch under its own name is fine.
    assert client.patch(f"/api/v1/branches/{first['id']}", headers=auth, json={"name": "North Mall", "code": "NM2"}).status_code == 200
    assert client.patch("/api/v1/branches/999999", headers=auth, json={"name": "Nope"}).status_code == 404
    assert client.delete(f"/api/v1/branches/{second['id']}", headers=auth).status_code == 200
    names = [row["name"] for row in client.get("/api/v1/branches", headers=auth).json()["items"]]
    assert "North Mall" in names and "South" not in names
